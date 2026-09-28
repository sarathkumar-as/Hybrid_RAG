import hashlib
import io
import json
import re
from dataclasses import dataclass

from pypdf import PdfReader
from docx import Document


@dataclass
class Chunk:
    id: str
    doc_id: str
    filename: str
    page: int
    text: str


def _split_text(text: str, doc_id: str, filename: str, page: int) -> list[Chunk]:
    # Keep headings and list boundaries: collapsing all whitespace makes a numbered
    # source list look like ordinary prose and invites an over-broad summary.
    text = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n"))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    chunks = []
    tokens = re.findall(r"\n|[^\s]+", text)
    for start in range(0, len(tokens), 170):
        section = " ".join(tokens[start:start + 210]).replace(" \n ", "\n").strip()
        if len(section) >= 25:
            chunks.append(Chunk(f"{doc_id}:{page}:{start}", doc_id, filename, page, section))
    return chunks


def document_chunks(data: bytes, filename: str, max_pages: int = 100, max_chunks: int = 150) -> list[Chunk]:
    suffix = filename.rsplit(".", 1)[-1].lower()
    doc_id = hashlib.sha256(data).hexdigest()
    if suffix not in ("pdf", "docx", "txt", "md"):
        raise ValueError("Supported formats: PDF, DOCX, TXT, Markdown (.md).")
    if suffix != "pdf":
        if suffix == "docx":
            document = Document(io.BytesIO(data))
            blocks = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    blocks.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(blocks)
        else:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise ValueError("TXT/Markdown must be UTF-8 encoded.") from exc
        # A non-PDF file has no reliable page numbers; section 1 denotes the document body.
        chunks = _split_text(text, doc_id, filename, 1)
    else:
        chunks = pdf_chunks(data, filename, max_pages, max_chunks)
    if len(chunks) > max_chunks:
        raise ValueError(f"Document exceeds the {max_chunks} chunk limit.")
    if not chunks:
        raise ValueError("No extractable text found. Run OCR on scanned PDFs first.")
    return chunks


def pdf_chunks(data: bytes, filename: str, max_pages: int = 100, max_chunks: int = 150) -> list[Chunk]:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are unsupported. Unlock the PDF before upload.")
    if len(reader.pages) > max_pages:
        raise ValueError(f"PDF exceeds the {max_pages} page limit.")
    doc_id = hashlib.sha256(data).hexdigest()
    chunks = []
    for page_no, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if not text:
            continue
        # Keep PDF page boundaries so page citations refer to one real page.
        chunks.extend(_split_text(text, doc_id, filename, page_no))
        if len(chunks) > max_chunks:
            raise ValueError(f"PDF exceeds the {max_chunks} chunk limit.")
    if not chunks:
        raise ValueError("No selectable text found. Run OCR on scanned PDFs first.")
    return chunks


def clean_entities(value: object) -> list[str]:
    if not isinstance(value, list):
        return []
    result = []
    for item in value:
        if not isinstance(item, str):
            continue
        name = re.sub(r"\s+", " ", item).strip().lower()[:80]
        if 2 <= len(name) <= 80 and name not in result:
            result.append(name)
    return result[:8]


def heading_query(question: str) -> str | None:
    """Recognize a short section title rather than treating it as a full question."""
    heading = question.strip().rstrip(":? ").strip()
    if not heading or len(heading) > 60 or len(heading.split()) > 4:
        return None
    if re.search(r"\b(what|which|who|where|when|why|how|list|explain|describe)\b", heading, re.I):
        return None
    return heading


def matching_heading_ids(question: str, ids: list[str], documents: list[str]) -> list[str]:
    """Prefer passages containing an actual heading, not an incidental word."""
    heading = heading_query(question)
    if not heading:
        return []
    pattern = re.compile(
        r"(?im)^[ \t]*(?:#{1,6}[ \t]*)?" + re.escape(heading)
        + r"[ \t]*(?::[ \t]*(?:\n|$)|\n|$)"
    )
    return [cid for cid, body in zip(ids, documents) if pattern.search(body)]


def citations_valid(answer: str, labels: set[str]) -> bool:
    """Require known citations on factual lines; permit Markdown section labels."""
    if not answer.strip():
        return False
    lines = [line.strip() for line in answer.splitlines() if line.strip()]
    if not lines:
        return False
    factual_lines = 0
    for line in lines:
        # A Markdown heading organizes cited bullets; it does not assert a
        # separate fact. Requiring a citation here rejected otherwise valid
        # answers whenever the model added a heading.
        if re.fullmatch(r"#{1,6}\s+[^\n]+", line):
            continue
        factual_lines += 1
        found = set(re.findall(r"\[S\d+\]", line))
        if not found or not found <= labels:
            return False
    return factual_lines > 0


def evidence_fallback(sources: list[dict], limit: int = 2) -> tuple[str, list[dict]]:
    """Show verbatim evidence when a generated answer cannot be verified."""
    selected = [source for source in sources if source.get("excerpt", "").strip()][:limit]
    if not selected:
        return "I could not find supporting text in the selected documents.", []
    lines = ["I could not verify a direct answer. These are the closest indexed passages:"]
    for source in selected:
        excerpt = " ".join(source["excerpt"].split())[:700]
        lines.append(f"> {excerpt} [{source['label']}]")
    return "\n\n".join(lines), selected


def fuse(vector_ids: list[str], graph_ids: list[str], limit: int = 6) -> list[tuple[str, float]]:
    scores: dict[str, float] = {}
    for weight, ids in ((1.0, vector_ids), (1.0, graph_ids)):
        for rank, cid in enumerate(dict.fromkeys(ids), 1):
            scores[cid] = scores.get(cid, 0) + weight / (rank + 60)
    return sorted(scores.items(), key=lambda item: (-item[1], item[0]))[:limit]


def parse_json(text: str) -> dict:
    try:
        result = json.loads(text)
        return result if isinstance(result, dict) else {}
    except (ValueError, TypeError):
        return {}
